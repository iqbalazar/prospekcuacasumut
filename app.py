import streamlit as st
import geopandas as gpd
import numpy as np
import pandas as pd
from scipy.spatial import cKDTree
import folium
from streamlit_folium import st_folium
import io
from PIL import Image
import shapely.geometry
from matplotlib.path import Path
import base64
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import re

# ==========================================
# KONFIGURASI DAN VARIABEL
# ==========================================
st.set_page_config(page_title="Sistem Prediksi Hujan Sumut", layout="wide")

FILE_SHP = "ADM_KOTKAB_SUMUT.shp"
FIELD_NAMA_KABUPATEN = "WADMKK"

# Bounding Box Gambar Georeference (Diperbarui sesuai input pengguna)
MIN_LON = 93.8232972577307009
MIN_LAT = -12.9974183620374362
MAX_LON = 143.7735419591668915
MAX_LAT = 10.1613314540829798

# Palet warna diperbarui sesuai permintaan user (RGB spesifik)
kategori_hujan = [
    {"kategori": "Hujan Ekstrem", "rgb": [255, 0, 255], "prioritas": 1, "hex": "#FF00FF"},
    {"kategori": "Hujan Sangat Lebat", "rgb": [255, 0, 0], "prioritas": 2, "hex": "#FF0000"},
    {"kategori": "Hujan Lebat", "rgb": [255, 155, 0], "prioritas": 3, "hex": "#FF9B00"},
    {"kategori": "Hujan Sedang", "rgb": [255, 255, 0], "prioritas": 4, "hex": "#FFFF00"},
    {"kategori": "Hujan Ringan", "rgb": [0, 255, 0], "prioritas": 5, "hex": "#00FF00"},
    {"kategori": "Cerah / Berawan", "rgb": [255, 255, 255], "prioritas": 6, "hex": "#FFFFFF"}
]

# Untuk Legenda Map dan Tabel
legenda_utama = {
    "Hujan Ekstrem": "#FF00FF",
    "Hujan Sangat Lebat": "#FF0000",
    "Hujan Lebat": "#FF9B00",
    "Hujan Sedang": "#FFFF00",
    "Hujan Ringan": "#00FF00",
    "Cerah / Berawan": "#FFFFFF"
}

warna_rgb = [k["rgb"] for k in kategori_hujan]
tree_warna = cKDTree(warna_rgb)

def coord_to_pixel(lon, lat, img_width, img_height):
    x = int((lon - MIN_LON) / (MAX_LON - MIN_LON) * img_width)
    y = int((MAX_LAT - lat) / (MAX_LAT - MIN_LAT) * img_height) 
    return x, y

def get_pixels_in_polygon(polygon, img_array):
    img_height, img_width, _ = img_array.shape
    minx, miny, maxx, maxy = polygon.bounds
    
    px_minx, px_maxy = coord_to_pixel(minx, miny, img_width, img_height) 
    px_maxx, px_miny = coord_to_pixel(maxx, maxy, img_width, img_height) 
    
    px_minx = max(0, px_minx)
    px_miny = max(0, px_miny)
    px_maxx = min(img_width, px_maxx)
    px_maxy = min(img_height, px_maxy)
    
    if px_minx >= px_maxx or px_miny >= px_maxy:
        return np.array([])
    
    x, y = np.meshgrid(np.arange(px_minx, px_maxx), np.arange(px_miny, px_maxy))
    x, y = x.flatten(), y.flatten()
    points = np.vstack((x, y)).T

    if isinstance(polygon, shapely.geometry.Polygon):
        polygons = [polygon]
    else: 
        polygons = polygon.geoms

    all_inside_pixels = []
    
    for poly in polygons:
        exterior_coords = np.array(poly.exterior.coords)
        pixel_coords = [coord_to_pixel(pt[0], pt[1], img_width, img_height) for pt in exterior_coords]
        
        path = Path(pixel_coords)
        mask = path.contains_points(points)
        
        points_inside = points[mask]
        
        for px, py in points_inside:
            all_inside_pixels.append(img_array[py, px, :3]) 

    return np.array(all_inside_pixels)


# ==========================================
# ANTARMUKA WEB (UI)
# ==========================================
st.title("🌦️ Sistem Deteksi Intensitas Curah Hujan Kab/Kota Sumut")
st.markdown("Unggah **satu atau beberapa** gambar prediksi cuaca BMKG untuk mengekstrak data intensitas hujan harian. Sistem akan otomatis membuat laporan rekap sepekan/beberapa hari jika Anda mengunggah banyak gambar sekaligus.")

# --- MEMORI SESI (AGAR HASIL TIDAK HILANG) ---
if 'df_laporan_gabungan' not in st.session_state:
    st.session_state.df_laporan_gabungan = None
if 'gdf_hasil_list' not in st.session_state:
    st.session_state.gdf_hasil_list = None
if 'uploaded_file_names' not in st.session_state:
    st.session_state.uploaded_file_names = []
if 'encoded_images' not in st.session_state:
    st.session_state.encoded_images = []

uploaded_files = st.file_uploader("Upload Gambar Prediksi Hujan (.jpg / .png)", type=["jpg", "jpeg", "png"], accept_multiple_files=True)

if uploaded_files:
    if st.button("Proses Intensitas Hujan", type="primary", use_container_width=True):
        with st.spinner("Sedang memindai koordinat, mengekstrak warna, dan memotong wilayah..."):
            
            gdf = gpd.read_file(FILE_SHP)
            if gdf.crs != "EPSG:4326":
                gdf = gdf.to_crs("EPSG:4326")
            gdf['geometry'] = gdf['geometry'].simplify(0.005) 

            semua_laporan = {}
            st.session_state.encoded_images = []
            st.session_state.uploaded_file_names = []
            st.session_state.gdf_hasil_list = {}
            
            progress_bar = st.progress(0)
            
            # Sort files by extracted date from filename
            def extract_date(filename):
                match = re.search(r'\d{8}', filename)
                return match.group(0) if match else filename
                
            uploaded_files_sorted = sorted(uploaded_files, key=lambda f: extract_date(f.name))

            total_tasks = len(uploaded_files_sorted) * len(gdf)
            task_count = 0

            # Proses setiap file gambar yang diunggah
            for idx, uploaded_file in enumerate(uploaded_files_sorted):
                nama_file = uploaded_file.name
                
                # Ekstrak tanggal dari nama file menggunakan Regex (mencari 8 digit berurutan)
                match = re.search(r'\d{8}', nama_file)
                if match:
                    header_kolom = match.group(0)
                else:
                    header_kolom = nama_file # Fallback jika format salah
                
                st.session_state.uploaded_file_names.append(header_kolom)
                
                # Simpan image untuk map overlay
                uploaded_file.seek(0)
                encoded_image = base64.b64encode(uploaded_file.read()).decode('utf-8')
                st.session_state.encoded_images.append('data:image/jpeg;base64,' + encoded_image)
                
                image = Image.open(uploaded_file).convert("RGB")
                img_array = np.array(image)
                
                laporan_harian = []
                
                for index, row in gdf.iterrows():
                    nama_wilayah = row[FIELD_NAMA_KABUPATEN]
                    
                    if pd.isna(nama_wilayah) or str(nama_wilayah).strip() == "":
                        task_count += 1
                        continue

                    polygon = row['geometry']
                    pixels_rgb = get_pixels_in_polygon(polygon, img_array)
                    
                    if len(pixels_rgb) == 0:
                        kategori_tertinggi = "Di Luar Peta"
                        hex_color = "#FFFFFF"
                    else:
                        distances, indices = tree_warna.query(pixels_rgb)
                        
                        # Toleransi ditingkatkan agar gradien tertangkap (150)
                        valid_indices = indices[distances < 150] 
                        
                        if len(valid_indices) == 0:
                            kategori_tertinggi = "Cerah / Berawan" 
                            hex_color = "#FFFFFF"
                        else:
                            prioritas_ditemukan = [kategori_hujan[i]["prioritas"] for i in valid_indices]
                            
                            # Filter: Jika ada warna SELAIN Cerah/Berawan (Prioritas 6), abaikan Cerah/Berawan sepenuhnya.
                            # Sekecil apapun warna hujan (Prioritas 1-5) yang masuk ke wilayah, catat yang tertinggi.
                            prioritas_hujan_saja = [p for p in prioritas_ditemukan if p < 6]
                            
                            if len(prioritas_hujan_saja) > 0:
                                prioritas_final = min(prioritas_hujan_saja)
                            else:
                                prioritas_final = 6 # Semuanya Cerah/Berawan
                                
                            kategori_terpilih = next(item for item in kategori_hujan if item["prioritas"] == prioritas_final)
                            kategori_tertinggi = kategori_terpilih["kategori"]
                            hex_color = kategori_terpilih["hex"]
                            
                    laporan_harian.append({
                        "Kabupaten_Kota": nama_wilayah,
                        f"Intensitas_{header_kolom}": kategori_tertinggi,
                        f"Warna_{header_kolom}": hex_color
                    })
                    
                    task_count += 1
                    progress_bar.progress(task_count / total_tasks)

                # Gabungkan laporan harian ke laporan utama
                df_harian = pd.DataFrame(laporan_harian)
                if len(semua_laporan) == 0:
                    semua_laporan = df_harian
                else:
                    semua_laporan = pd.merge(semua_laporan, df_harian, on="Kabupaten_Kota")
                
                # Simpan gdf harian untuk peta
                st.session_state.gdf_hasil_list[header_kolom] = gdf.merge(df_harian, left_on=FIELD_NAMA_KABUPATEN, right_on="Kabupaten_Kota")

            st.session_state.df_laporan_gabungan = semua_laporan

# --- MENAMPILKAN HASIL DARI MEMORI SESI ---
if st.session_state.df_laporan_gabungan is not None:
    st.success(f"✅ Pemrosesan {len(st.session_state.uploaded_file_names)} Gambar Selesai!")
    st.divider()
    
    # === BAGIAN TABEL & EXCEL ===
    st.subheader("📋 Tabel Hasil Rekapitulasi Curah Hujan")
    
    df_tampil = st.session_state.df_laporan_gabungan.copy()
    
    # Buat fungsi pewarnaan cell dataframe Streamlit
    def color_cells(val):
        for k, v in legenda_utama.items():
            if val == k:
                color = "black" if v in ["#FFFF00", "#00FF00", "#FFFFFF"] else "white"
                # Add border so white is visible
                border = "border: 1px solid #ccc;" if v == "#FFFFFF" else ""
                return f'background-color: {v}; color: {color}; font-weight: bold; {border}'
        return ''
    
    # Ambil kolom yang hanya berisi 'Intensitas' (Hapus kolom warna Hex untuk tampilan web)
    kolom_intensitas = [col for col in df_tampil.columns if col.startswith("Intensitas_")]
    df_tampil_bersih = df_tampil[["Kabupaten_Kota"] + kolom_intensitas]
    
    # Terapkan gaya ke tabel web
    st.dataframe(
        df_tampil_bersih.style.map(color_cells, subset=kolom_intensitas), 
        width='stretch', 
        height=400
    )
    
    # Proses Excel Custom (Mirip Format Lampiran User)
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap Hujan"
    
    # Header
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2+len(kolom_intensitas))
    ws['A1'] = "POTENSI HUJAN KAB./KOTA DI SUMATERA UTARA"
    ws['A1'].font = Font(bold=True)
    ws['A1'].alignment = Alignment(horizontal="center")
    
    # Kolom Header
    ws['A3'] = "No"
    ws['B3'] = "Provinsi"
    ws.column_dimensions['B'].width = 30
    
    for i, col in enumerate(kolom_intensitas):
        cell = ws.cell(row=3, column=3+i)
        # Menghapus kata "Intensitas_" agar hanya menyisakan tanggal di header Excel
        cell.value = col.replace("Intensitas_", "") 
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[openpyxl.utils.get_column_letter(3+i)].width = 15

    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Isi Data & Warna Excel
    for r_idx, row in df_tampil.iterrows():
        ws.cell(row=r_idx+4, column=1, value=r_idx+1).border = thin_border
        ws.cell(row=r_idx+4, column=2, value=row["Kabupaten_Kota"]).border = thin_border
        
        for c_idx, col in enumerate(kolom_intensitas):
            val = row[col]
            warna_hex = row[f"Warna_{col.replace('Intensitas_', '')}"].replace("#", "")
            
            cell = ws.cell(row=r_idx+4, column=3+c_idx)
            cell.value = "" # Kosongkan teks agar full warna seperti lampiran
            cell.fill = PatternFill(start_color=warna_hex, end_color=warna_hex, fill_type="solid")
            cell.border = thin_border

    # Tambahkan Legenda di bawah tabel Excel
    start_leg_row = len(df_tampil) + 6
    for i, (kategori, hex_color) in enumerate(legenda_utama.items()):
        ws.cell(row=start_leg_row+i, column=1).fill = PatternFill(start_color=hex_color.replace("#", ""), end_color=hex_color.replace("#", ""), fill_type="solid")
        ws.cell(row=start_leg_row+i, column=2, value=kategori)

    wb.save(output)
    excel_data = output.getvalue()
    
    st.download_button(
        label="📥 Download Laporan Matriks (Excel)",
        data=excel_data,
        file_name="Laporan_Hujan_Rekap_Sumut.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )
    
    st.divider()

    # === BAGIAN PETA ===
    st.subheader("🗺️ Peta Visualisasi Hasil Overlay")
    
    # Pilih Hari/Gambar mana yang ingin dilihat di Peta
    selected_map_file = st.selectbox("Pilih Peta Hari / Gambar yang ingin ditampilkan:", st.session_state.uploaded_file_names)
    
    col_map1, col_map2 = st.columns([1, 4])
    
    with col_map1:
        st.markdown("**Kontrol Tampilan Peta:**")
        opacity_gambar = st.slider("Opasitas Gambar Prediksi", 0.0, 1.0, 0.6, 0.1)
        opacity_poligon = st.slider("Opasitas Poligon Hasil", 0.0, 1.0, 0.7, 0.1)
        
        st.markdown("**Legenda BMKG:**")
        for k, v in legenda_utama.items():
            st.markdown(f'<div style="display:flex;align-items:center;margin-bottom:5px;"><div style="width:20px;height:20px;background-color:{v};margin-right:10px;border:1px solid #ccc;"></div><span style="font-size:14px;">{k}</span></div>', unsafe_allow_html=True)
            
    with col_map2:
        # Peta menggunakan plugins Fullscreen agar bisa dimaximize
        import folium.plugins as plugins
        
        peta = folium.Map(location=[2.5, 99.0], zoom_start=6, tiles="OpenStreetMap")
        plugins.Fullscreen(position='topright', title='Full Screen', title_cancel='Exit Full Screen', force_separate_button=True).add_to(peta)
        
        # Cari index gambar yang dipilih
        idx_selected = st.session_state.uploaded_file_names.index(selected_map_file)
        image_url = st.session_state.encoded_images[idx_selected]
        gdf_hari_ini = st.session_state.gdf_hasil_list[selected_map_file]
        
        bounds = [[MIN_LAT, MIN_LON], [MAX_LAT, MAX_LON]]
        folium.raster_layers.ImageOverlay(
            image=image_url,
            bounds=bounds,
            opacity=opacity_gambar,
            name=f"Gambar: {selected_map_file}",
            interactive=True
        ).add_to(peta)

        folium.GeoJson(
            gdf_hari_ini,
            name="Hasil Ekstraksi Wilayah",
            style_function=lambda feature: {
                'fillColor': feature['properties'][f'Warna_{selected_map_file}'],
                'color': 'black',
                'weight': 1,
                'fillOpacity': opacity_poligon
            },
            tooltip=folium.features.GeoJsonTooltip(
                fields=[FIELD_NAMA_KABUPATEN, f'Intensitas_{selected_map_file}'],
                aliases=['Wilayah:', 'Intensitas:'],
                localize=True
            )
        ).add_to(peta)
        
        folium.LayerControl().add_to(peta)
        
        st_folium(peta, width="100%", height=600)